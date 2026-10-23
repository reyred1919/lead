"""Export generation. Row scope = visible_interactions(user) + filters; field scope = PII masking."""
import csv
import io
import json
import secrets
import zipfile
from collections import defaultdict
from datetime import timedelta
from pathlib import Path

from django.conf import settings
from django.db.models import Prefetch
from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo

from apps.common.audit import audit
from apps.common.spreadsheet import safe_cell
from apps.common.text import mask_email, mask_phone
from apps.crm.models import (
    Activity,
    Contact,
    Interaction,
    InteractionProduct,
    Opportunity,
    OpportunityTag,
    Referral,
    Task,
)
from apps.crm.scoping import visible_interactions
from apps.crm.views import apply_lead_filters
from apps.iam.authz import ALL_PRODUCTS, authz_for

from .models import DataExport
from .schema import DATASETS, DEFAULT_DATASETS, SCHEMA_VERSION, columns


def _iso(dt):
    return dt.isoformat() if dt else None


def _s(v):
    return str(v) if v else None


class _Params(dict):
    def get(self, key, default=None):
        value = super().get(key, default)
        if isinstance(value, list):
            return ",".join(str(v) for v in value)
        return value


def collect(user, filters, datasets):
    base = visible_interactions(user, Interaction.objects.all())
    qs = apply_lead_filters(base, _Params(filters or {}), user).select_related(
        "event", "contact", "account", "created_by", "opportunity__owner"
    ).prefetch_related(
        Prefetch("product_links", queryset=InteractionProduct.objects.select_related("product")),
        Prefetch("opportunity__tag_links", queryset=OpportunityTag.objects.select_related("tag")),
    ).order_by("captured_at")
    interactions = list(qs)
    ids = [i.id for i in interactions]

    az = authz_for(user)
    pii_all = az.has_global("pii.view")
    pii_events = {g.event_id for g in az.event_grants("pii.view") if g.product_scope is ALL_PRODUCTS}

    def pii_ok(i):
        opp = getattr(i, "_opp", None)
        return pii_all or i.event_id in pii_events or i.created_by_id == user.id or i.triage_owner_id == user.id or bool(opp and opp.owner_id == user.id)

    opps = {}
    for i in interactions:
        try:
            i._opp = i.opportunity
        except Opportunity.DoesNotExist:
            i._opp = None
        if i._opp:
            opps[i._opp.id] = i._opp
    pii_contacts = {i.contact_id for i in interactions if pii_ok(i)}

    out = {}
    want = set(datasets)

    if "fact_leads" in want:
        tasks_by = defaultdict(lambda: [0, 0])
        for t in Task.objects.filter(interaction_id__in=ids).values("interaction_id", "status"):
            if t["status"] == "OPEN":
                tasks_by[t["interaction_id"]][0] += 1
            elif t["status"] == "DONE":
                tasks_by[t["interaction_id"]][1] += 1
        latest_ref = {}
        for r in Referral.objects.filter(interaction_id__in=ids).order_by("created_at").values("interaction_id", "status"):
            latest_ref[r["interaction_id"]] = r["status"]
        contact_events = defaultdict(set)
        for i in interactions:
            contact_events[i.contact_id].add(i.event_id)
        rows = []
        for i in interactions:
            opp = i._opp
            links = list(i.product_links.all())
            tags = [l.tag for l in opp.tag_links.all()] if opp else []
            pii = pii_ok(i)
            rows.append({
                "interaction_id": str(i.id), "event_id": str(i.event_id), "event_code": i.event_code, "event_name": i.event.display_name,
                "captured_at": _iso(i.captured_at), "captured_date": i.captured_at.date().isoformat(),
                "created_by_id": str(i.created_by_id), "created_by_name": i.created_by.full_name,
                "contact_id": str(i.contact_id), "contact_name": i.contact.full_name,
                "contact_mobile": i.contact.mobile if pii else mask_phone(i.contact.mobile),
                "account_id": _s(i.account_id), "account_name": i.account.name if i.account_id else None,
                "product_codes": "|".join(l.product_code for l in links), "product_names": "|".join(l.product.name_fa for l in links),
                "product_count": len(links), "quality": i.quality, "interaction_status": i.status,
                "opportunity_id": _s(opp.id) if opp else None, "opportunity_stage": opp.stage if opp else None,
                "owner_id": _s(opp.owner_id) if opp else _s(i.triage_owner_id),
                "owner_name": (opp.owner.full_name if opp and opp.owner_id else None),
                "tag_codes": "|".join(t.code for t in tags), "tag_labels": "|".join(t.name_fa for t in tags),
                "open_tasks": tasks_by[i.id][0], "done_tasks": tasks_by[i.id][1],
                "latest_referral_status": latest_ref.get(i.id), "contact_event_count": len(contact_events[i.contact_id]),
            })
        out["fact_leads"] = rows

    if "interactions" in want:
        out["interactions"] = [{
            "interaction_id": str(i.id), "event_id": str(i.event_id), "event_code": i.event_code, "created_by_id": str(i.created_by_id),
            "created_by_name": i.created_by.full_name, "captured_at": _iso(i.captured_at), "captured_offline": i.captured_offline,
            "client_record_id": str(i.client_record_id), "contact_id": str(i.contact_id), "account_id": _s(i.account_id),
            "captured_contact_name": i.captured_contact_name, "captured_org_name": i.captured_org_name, "quality": i.quality,
            "status": i.status, "triage_owner_id": _s(i.triage_owner_id), "note": i.note, "updated_at": _iso(i.updated_at),
        } for i in interactions]

    if "interaction_products" in want:
        out["interaction_products"] = [
            {"interaction_id": str(i.id), "product_id": str(l.product_id), "product_code": l.product_code}
            for i in interactions for l in i.product_links.all()
        ]

    if "accounts" in want:
        account_ids = {i.account_id for i in interactions if i.account_id}
        from apps.crm.models import Account

        out["accounts"] = [{
            "account_id": str(a.id), "name": a.name, "official_name": a.official_name, "alt_name": a.alt_name, "industry": a.industry,
            "org_type": a.org_type, "city": a.city, "website": a.website, "legal_id": a.legal_id, "status": a.status,
            "merged_into_id": _s(a.merged_into_id), "created_at": _iso(a.created_at), "updated_at": _iso(a.updated_at),
        } for a in Account.objects.filter(pk__in=account_ids).order_by("name")]

    if "contacts" in want:
        contact_ids = {i.contact_id for i in interactions}
        rows = []
        for c in Contact.objects.filter(pk__in=contact_ids).order_by("full_name"):
            pii = c.id in pii_contacts
            rows.append({
                "contact_id": str(c.id), "full_name": c.full_name, "title": c.title,
                "mobile": c.mobile if pii else mask_phone(c.mobile), "email": c.email if pii else mask_email(c.email),
                "alt_phone": c.alt_phone if pii else mask_phone(c.alt_phone), "primary_account_id": _s(c.primary_account_id),
                "status": c.status, "merged_into_id": _s(c.merged_into_id), "created_at": _iso(c.created_at), "updated_at": _iso(c.updated_at),
            })
        out["contacts"] = rows

    if "opportunities" in want:
        out["opportunities"] = [{
            "opportunity_id": str(o.id), "source_interaction_id": str(o.source_interaction_id), "origin_event_id": str(o.origin_event_id),
            "account_id": _s(o.account_id), "contact_id": str(o.contact_id), "stage": o.stage, "owner_id": _s(o.owner_id),
            "owner_name": o.owner.full_name if o.owner_id else None, "qualified_by_id": _s(o.qualified_by_id),
            "qualified_at": _iso(o.qualified_at), "urgency": o.urgency or None, "business_problem": o.business_problem,
            "use_case": o.use_case, "next_step": o.next_step, "next_step_due_at": _iso(o.next_step_due_at), "closed_at": _iso(o.closed_at),
            "tag_codes": "|".join(l.tag.code for l in o.tag_links.all()), "updated_at": _iso(o.updated_at),
        } for o in opps.values()]

    if "opportunity_tags" in want:
        out["opportunity_tags"] = [{
            "opportunity_id": str(l.opportunity_id), "tag_id": str(l.tag_id), "tag_code": l.tag.code, "tag_name_fa": l.tag.name_fa,
            "tag_color": l.tag.color, "tag_status": l.tag.status, "applied_by_id": _s(l.created_by_id), "applied_at": _iso(l.created_at),
        } for l in OpportunityTag.objects.filter(opportunity_id__in=list(opps)).select_related("tag")]

    if "referrals" in want:
        out["referrals"] = [{
            "referral_id": str(r.id), "interaction_id": str(r.interaction_id), "opportunity_id": _s(r.opportunity_id),
            "event_id": str(r.event_id), "purpose": r.purpose, "from_user_id": str(r.from_user_id), "to_user_id": str(r.to_user_id),
            "status": r.status, "reason": r.reason, "response_note": r.response_note, "due_at": _iso(r.due_at),
            "responded_at": _iso(r.responded_at), "created_at": _iso(r.created_at),
        } for r in Referral.objects.filter(interaction_id__in=ids).order_by("created_at")]

    if "tasks" in want:
        out["tasks"] = [{
            "task_id": str(t.id), "interaction_id": str(t.interaction_id), "opportunity_id": _s(t.opportunity_id), "title": t.title,
            "owner_id": str(t.owner_id), "status": t.status, "due_at": _iso(t.due_at), "completed_at": _iso(t.completed_at),
            "outcome": t.outcome, "created_at": _iso(t.created_at),
        } for t in Task.objects.filter(interaction_id__in=ids).order_by("created_at")]

    if "activities" in want:
        out["activities"] = [{
            "activity_id": str(a.id), "interaction_id": str(a.interaction_id), "opportunity_id": _s(a.opportunity_id), "type": a.type,
            "actor_id": _s(a.actor_id), "notes": a.notes, "created_at": _iso(a.created_at),
        } for a in Activity.objects.filter(interaction_id__in=ids).order_by("created_at")]

    return out, (not pii_all)


def _metadata(job, user, rows):
    return {
        "schema_version": SCHEMA_VERSION,
        "export_id": str(job.id),
        "generated_at": timezone.now().isoformat(),
        "generated_by": {"id": str(user.id), "name": user.full_name},
        "filters": job.filters,
        "datasets": {k: len(v) for k, v in rows.items()},
        "timestamps": "UTC ISO-8601",
        "pii_masked_where_unauthorized": True,
        "formula_injection_defense": "Text cells starting with = + - @ tab CR LF | % are prefixed with an apostrophe (CSV/XLSX only).",
        "demo_data": settings.DEMO_MODE,
        "csv_encoding": "UTF-8 with BOM (Excel-compatible)",
    }


def _write_csv(rows, dataset):
    buf = io.StringIO()
    buf.write("﻿")
    writer = csv.writer(buf, quoting=csv.QUOTE_MINIMAL, lineterminator="\r\n")
    cols = columns(dataset)
    writer.writerow(cols)
    for row in rows:
        writer.writerow([safe_cell(row.get(c)) if row.get(c) is not None else "" for c in cols])
    return buf.getvalue().encode("utf-8")


def _write_xlsx(rows_by_ds, meta):
    wb = Workbook()
    ws = wb.active
    ws.title = "metadata"
    ws.sheet_view.rightToLeft = True
    header = Font(bold=True, color="FFFFFF")
    fill = PatternFill("solid", fgColor="10294C")
    for k, v in meta.items():
        ws.append([k, safe_cell(json.dumps(v, ensure_ascii=False) if isinstance(v, (dict, list)) else (str(v) if v is not None else ""))])
    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 90
    for ds, rows in rows_by_ds.items():
        sheet = wb.create_sheet(ds[:31])
        sheet.sheet_view.rightToLeft = True
        cols = columns(ds)
        sheet.append(cols)
        for cell in sheet[1]:
            cell.font, cell.fill = header, fill
        for row in rows:
            values = []
            for c in cols:
                v = row.get(c)
                values.append(safe_cell(v) if isinstance(v, str) else v)
            sheet.append(values)
        for idx, c in enumerate(cols, start=1):
            sheet.column_dimensions[sheet.cell(row=1, column=idx).column_letter].width = min(40, max(12, len(c) + 4))
        sheet.freeze_panes = "A2"
        if rows:
            ref = f"A1:{sheet.cell(row=len(rows) + 1, column=len(cols)).coordinate}"
            table = Table(displayName=f"t_{ds}"[:250], ref=ref)
            table.tableStyleInfo = TableStyleInfo(name="TableStyleLight9", showRowStripes=True)
            sheet.add_table(table)
    dic = wb.create_sheet("dictionary")
    dic.sheet_view.rightToLeft = True
    dic.append(["dataset", "column", "description", "privacy_class"])
    for cell in dic[1]:
        cell.font, cell.fill = header, fill
    for ds in rows_by_ds:
        for col, desc, cls in DATASETS[ds]:
            dic.append([ds, col, desc, cls])
    # Ensure no cell is stored as a formula, even if a value slipped past safe_cell.
    for sheet in wb.worksheets:
        for row in sheet.iter_rows():
            for cell in row:
                if cell.data_type == "f":
                    cell.data_type = "s"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def render(job, user):
    datasets = [d for d in (job.datasets or DEFAULT_DATASETS) if d in DATASETS] or DEFAULT_DATASETS
    rows, masked = collect(user, job.filters, datasets)
    meta = _metadata(job, user, rows)
    stamp = timezone.now().strftime("%Y%m%d-%H%M")
    if job.format == "JSON":
        content = json.dumps({"metadata": meta, "dictionary": {d: DATASETS[d] for d in rows}, "datasets": rows}, ensure_ascii=False, indent=1).encode()
        name = f"expo-leads-{stamp}.json"
    elif job.format == "CSV":
        if len(rows) == 1:
            ds = next(iter(rows))
            content, name = _write_csv(rows[ds], ds), f"expo-leads-{ds}-{stamp}.csv"
        else:
            buf = io.BytesIO()
            with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zf:
                for ds, r in rows.items():
                    zf.writestr(f"{ds}.csv", _write_csv(r, ds))
                zf.writestr("metadata.json", json.dumps({**meta, "dictionary": {d: DATASETS[d] for d in rows}}, ensure_ascii=False, indent=1))
            content, name = buf.getvalue(), f"expo-leads-{stamp}.zip"
    else:
        content, name = _write_xlsx(rows, meta), f"expo-leads-{stamp}.xlsx"
    return content, name, {k: len(v) for k, v in rows.items()}, masked


def run_export(job_id):
    job = DataExport.objects.select_related("created_by").get(pk=job_id)
    job.status = "RUNNING"
    job.save(update_fields=["status"])
    try:
        content, name, counts, masked = render(job, job.created_by)
        settings.EXPORT_ROOT.mkdir(parents=True, exist_ok=True)
        key = secrets.token_hex(16)
        Path(settings.EXPORT_ROOT, key).write_bytes(content)
        job.file_key, job.file_name, job.file_size = key, name, len(content)
        job.row_counts, job.pii_masked = counts, masked
        job.status = "READY"
        job.completed_at = timezone.now()
        job.expires_at = timezone.now() + timedelta(hours=settings.EXPORT_TTL_HOURS)
        job.save()
        audit(job.created_by, "export.generated", "data_export", job.id, metadata={"format": job.format, "rows": counts, "filters": job.filters, "pii_masked": masked})
    except Exception as exc:  # noqa: BLE001
        job.status = "FAILED"
        job.error = str(exc)[:500]
        job.save()
        raise
    return job


def purge_expired():
    count = 0
    for job in DataExport.objects.filter(status="READY", expires_at__lt=timezone.now()):
        if job.file_key:
            Path(settings.EXPORT_ROOT, job.file_key).unlink(missing_ok=True)
        job.status = "EXPIRED"
        job.file_key = ""
        job.save()
        count += 1
    return count
