"""Exports (§16, §22.6), tags (§23.2–23.5), branding upload safety (§23.1)."""
import csv
import io
import json
import zipfile

import pytest
from django.core.files.uploadedfile import SimpleUploadedFile
from openpyxl import load_workbook
from PIL import Image

from apps.common.models import AuditLog
from apps.common.spreadsheet import safe_cell
from apps.crm.models import Tag
from apps.reports.models import DataExport

from .conftest import client_for

pytestmark = pytest.mark.django_db


def export(user, fmt="JSON", **body):
    c = client_for(user)
    resp = c.post("/api/v1/exports", {"format": fmt, **body}, format="json")
    return c, resp


def download(c, resp):
    r = c.get(resp.data["download_url"])
    assert r.status_code == 200, r
    return b"".join(r.streaming_content)


@pytest.mark.parametrize("value,expected", [
    ("=HYPERLINK(\"x\")", "'=HYPERLINK(\"x\")"), ("+98912", "'+98912"), ("-1", "'-1"), ("@SUM", "'@SUM"),
    ("\t=1", "'\t=1"), ("normal", "normal"), ("علی", "علی"), (5, 5), (None, None),
])
def test_safe_cell(value, expected):
    assert safe_cell(value) == expected


def test_export_formula_injection_defended_in_csv_and_xlsx(world, capture):
    capture(world.specialist, contact_name="=HYPERLINK(\"http://evil\",\"x\")", note="@SUM(A1:A9)", org_name="+cmd|' /C calc'!A0")
    c, resp = export(world.admin, "CSV", datasets=["fact_leads"])
    content = download(c, resp).decode("utf-8")
    assert content.startswith("﻿")
    rows = list(csv.reader(io.StringIO(content.lstrip("﻿"))))
    header, row = rows[0], rows[1]
    assert row[header.index("contact_name")].startswith("'=")
    assert row[header.index("account_name")].startswith("'+")
    c, resp = export(world.admin, "XLSX", datasets=["fact_leads", "interactions"])
    wb = load_workbook(io.BytesIO(download(c, resp)))
    for ws in wb.worksheets:
        for r in ws.iter_rows():
            for cell in r:
                assert cell.data_type != "f"
                if isinstance(cell.value, str) and cell.value:
                    assert cell.value[0] not in "=+-@" or ws.title in ("dictionary",)
    sheet = wb["interactions"]
    head = [c.value for c in sheet[1]]
    assert sheet.cell(row=2, column=head.index("note") + 1).value == "'@SUM(A1:A9)"
    assert {"metadata", "dictionary", "fact_leads", "interactions"} <= set(wb.sheetnames)


def test_export_is_scoped_and_masks_pii(world, capture):
    capture(world.specialist)
    capture(world.spec_f, event=world.f, products=[world.z], contact_name="دیگر", mobile="09125550000")
    c, resp = export(world.planmgr)
    data = json.loads(download(c, resp))
    assert data["metadata"]["schema_version"] == "1.0"
    assert {r["event_code"] for r in data["datasets"]["fact_leads"]} == {"E"}
    assert all(r["contact_id"] for r in data["datasets"]["contacts"])
    assert len(data["datasets"]["contacts"]) == 1
    c, resp = export(world.dg)
    data = json.loads(download(c, resp))
    assert {r["event_code"] for r in data["datasets"]["fact_leads"]} == {"E", "F"}
    assert all("•" not in r["contact_mobile"] for r in data["datasets"]["fact_leads"])


def test_export_event_filter_outside_scope_denied(world, capture):
    c, resp = export(world.planmgr, filters={"event_id": [str(world.f.id)]})
    assert resp.status_code == 403
    _, resp = export(world.specialist)
    assert resp.status_code == 403  # booth specialist has no export.run


def test_export_download_requires_owner_and_token(world, capture):
    capture(world.specialist)
    c, resp = export(world.admin, "CSV", datasets=["accounts", "contacts"])
    job = DataExport.objects.get(pk=resp.data["id"])
    other = client_for(world.dg)
    assert other.get(resp.data["download_url"]).status_code == 404
    assert c.get(f"/api/v1/exports/{job.id}/download?token=wrong").status_code == 404
    z = zipfile.ZipFile(io.BytesIO(download(c, resp)))
    assert {"accounts.csv", "contacts.csv", "metadata.json"} == set(z.namelist())
    assert AuditLog.objects.filter(action="export.downloaded", entity_id=str(job.id)).exists()
    assert AuditLog.objects.filter(action="export.download_denied").exists()
    from django.utils import timezone

    DataExport.objects.filter(pk=job.pk).update(expires_at=timezone.now())
    assert c.get(resp.data["download_url"]).status_code == 404


def test_cross_event_contact_history_in_export(world, capture):
    from .conftest import member

    member(world.f, world.specialist)
    capture(world.specialist)
    capture(world.specialist, event=world.f, products=[world.z])
    c, resp = export(world.mdm, filters={"event_id": [str(world.e.id), str(world.f.id)]})
    facts = json.loads(download(c, resp))["datasets"]["fact_leads"]
    assert len(facts) == 2 and all(f["contact_event_count"] == 2 for f in facts)


# --------------------------------------------------------------- Tags ---

def qualified(world, capture):
    lid = capture(world.specialist).data["id"]
    data = client_for(world.planmgr).post(f"/api/v1/interactions/{lid}/qualify", {}, format="json").data
    return lid, data["opportunity"]


def test_tag_taxonomy_crud_and_archive(world):
    admin, mdm = client_for(world.admin), client_for(world.mdm)
    t = mdm.post("/api/v1/tags", {"name_fa": "پیگیری فوری", "name_en": "Urgent", "color": "rose"}, format="json")
    assert t.status_code == 201
    dup = mdm.post("/api/v1/tags", {"name_fa": "پيگيري  فوري", "color": "blue"}, format="json")
    assert dup.status_code == 400  # Persian/Arabic-aware duplicate prevention
    assert mdm.post("/api/v1/tags", {"name_fa": "x", "color": "#ff00ff"}, format="json").status_code == 400
    edited = admin.patch(f"/api/v1/tags/{t.data['id']}", {"color": "violet", "version": t.data["version"]}, format="json")
    assert edited.data["color"] == "violet"
    assert client_for(world.planmgr).post("/api/v1/tags", {"name_fa": "y"}, format="json").status_code == 403
    archived = mdm.post(f"/api/v1/tags/{t.data['id']}/archive", {"version": edited.data["version"]}, format="json")
    assert archived.data["status"] == "ARCHIVED"
    active = client_for(world.specialist).get("/api/v1/tags?active=1").data
    assert t.data["id"] not in [x["id"] for x in active]


def test_scenario_k_l_m_tag_permissions_and_history(world, capture):
    lid, opp = qualified(world, capture)
    urgent = Tag.objects.create(code="u", name_fa="پیگیری فوری", name_normalized="پیگیری فوری", color="rose")
    demo = Tag.objects.create(code="d", name_fa="نیازمند دمو", name_normalized="نیازمند دمو", color="blue")
    old = Tag.objects.create(code="o", name_fa="قدیمی", name_normalized="قدیمی", color="slate")
    pm = client_for(world.planmgr)
    assert pm.post(f"/api/v1/opportunities/{opp['id']}/tags", {"tag_id": str(demo.id)}, format="json").status_code == 201
    assert pm.post(f"/api/v1/opportunities/{opp['id']}/tags", {"tag_id": str(old.id)}, format="json").status_code == 201
    assert pm.post(f"/api/v1/opportunities/{opp['id']}/tags", {"tag_id": str(urgent.id)}, format="json").status_code == 201
    assert pm.delete(f"/api/v1/opportunities/{opp['id']}/tags/{demo.id}").status_code == 200
    # Archived tag persists on the record but can no longer be applied.
    old.status = "ARCHIVED"
    old.save()
    detail = pm.get(f"/api/v1/interactions/{lid}").data
    assert {t["code"] for t in detail["opportunity"]["tags"]} == {"u", "o"}
    assert pm.post(f"/api/v1/opportunities/{opp['id']}/tags", {"tag_id": str(demo.id)}, format="json").status_code == 201
    pm.delete(f"/api/v1/opportunities/{opp['id']}/tags/{demo.id}")
    lid2, opp2 = qualified(world, capture)
    assert pm.post(f"/api/v1/opportunities/{opp2['id']}/tags", {"tag_id": str(old.id)}, format="json").status_code == 422
    types = [a["type"] for a in pm.get(f"/api/v1/interactions/{lid}/activities").data]
    assert types.count("TAG_APPLIED") == 4 and types.count("TAG_REMOVED") == 2
    assert AuditLog.objects.filter(action="TAG_APPLIED").count() == 4
    # L: specialist (creator, no detail edit) cannot tag via direct API.
    spec = client_for(world.specialist)
    assert spec.post(f"/api/v1/opportunities/{opp['id']}/tags", {"tag_id": str(demo.id)}, format="json").status_code == 403
    assert spec.delete(f"/api/v1/opportunities/{opp['id']}/tags/{urgent.id}").status_code == 403
    assert spec.get(f"/api/v1/interactions/{lid}").data["permissions"]["edit_detail"] is False
    # M: executive viewer sees tags read-only.
    ex = client_for(world.exec)
    assert {t["code"] for t in ex.get(f"/api/v1/interactions/{lid}").data["opportunity"]["tags"]} == {"u", "o"}
    assert ex.post(f"/api/v1/opportunities/{opp['id']}/tags", {"tag_id": str(demo.id)}, format="json").status_code == 403
    # Tag filter honours scope.
    assert {r["id"] for r in pm.get(f"/api/v1/interactions?tag_id={urgent.id}").data["results"]} == {lid}
    assert client_for(world.spec_f).get(f"/api/v1/interactions?tag_id={urgent.id}").data["results"] == []
    # N: exports carry tag ids/codes/labels/colors.
    c, resp = export(world.planmgr, datasets=["opportunity_tags", "fact_leads"])
    data = json.loads(download(c, resp))
    row = data["datasets"]["opportunity_tags"][0]
    assert {"tag_id", "tag_code", "tag_name_fa", "tag_color"} <= set(row)
    assert any("u" in f["tag_codes"].split("|") for f in data["datasets"]["fact_leads"])


def test_unqualified_lead_cannot_be_tagged_without_opportunity(world, capture):
    lid = capture(world.specialist).data["id"]
    detail = client_for(world.planmgr).get(f"/api/v1/interactions/{lid}").data
    assert detail["opportunity"] is None


# ----------------------------------------------------------- Branding ---

def png_bytes(size=(200, 60), fmt="PNG"):
    buf = io.BytesIO()
    Image.new("RGBA" if fmt != "JPEG" else "RGB", size, (16, 41, 76, 255) if fmt != "JPEG" else (16, 41, 76)).save(buf, format=fmt)
    return buf.getvalue()


def upload(user, content, name="logo.png", variant="LIGHT", ctype="image/png"):
    return client_for(user).post("/api/v1/branding/logo", {"variant": variant, "file": SimpleUploadedFile(name, content, content_type=ctype)}, format="multipart")


def test_scenario_i_logo_upload_replace_reset(world):
    from rest_framework.test import APIClient

    public = APIClient().get("/api/v1/branding").data
    assert public["placeholder"] is True
    r1 = upload(world.admin, png_bytes())
    assert r1.status_code == 201 and r1.data["light"]["width"] == 200
    first_id = r1.data["light"]["id"]
    r2 = upload(world.mdm, png_bytes((300, 90), "WEBP"), name="../../etc/passwd.webp", ctype="image/webp")
    assert r2.status_code == 201 and r2.data["revision"] > r1.data["revision"]
    from apps.branding.models import BrandAsset

    old = BrandAsset.objects.get(pk=first_id)
    assert old.is_active is False and old.archived_at
    new = BrandAsset.objects.get(pk=r2.data["light"]["id"])
    assert "/" not in new.storage_key and ".." not in new.storage_key and new.mime == "image/png"
    served = APIClient().get(r2.data["light"]["url"])
    assert served.status_code == 200 and served["Content-Type"] == "image/png"
    reset = client_for(world.admin).delete("/api/v1/branding/logo/LIGHT")
    assert reset.data["placeholder"] is True
    assert AuditLog.objects.filter(action="branding.logo_uploaded").count() == 2
    assert AuditLog.objects.filter(action="branding.logo_reset").exists()


@pytest.mark.parametrize("content,name", [
    (b'<svg xmlns="http://www.w3.org/2000/svg"><script>alert(1)</script></svg>', "logo.svg"),
    (b"<?xml version='1.0'?><svg></svg>", "logo.png"),
    (b"MZ\x90\x00 not an image", "logo.png"),
    (b"\x89PNG\r\n\x1a\n" + b"garbage" * 10, "logo.png"),
])
def test_malicious_or_invalid_uploads_rejected_and_audited(world, content, name):
    resp = upload(world.admin, content, name=name)
    assert resp.status_code == 400
    assert AuditLog.objects.filter(action="branding.logo_rejected").exists()


def test_upload_limits_and_permissions(world, settings):
    assert upload(world.admin, png_bytes((10, 10))).status_code == 400  # too small
    settings.BRAND_MAX_UPLOAD_BYTES = 100
    assert upload(world.admin, png_bytes()).status_code == 400
    settings.BRAND_MAX_UPLOAD_BYTES = 2 * 1024 * 1024
    assert upload(world.dg, png_bytes()).status_code == 403
    assert upload(world.specialist, png_bytes()).status_code == 403


def test_polyglot_payload_is_stripped_by_reencoding(world):
    payload = png_bytes() + b"<?php system($_GET['c']); ?>"
    resp = upload(world.admin, payload)
    assert resp.status_code == 201
    from apps.branding.models import BrandAsset
    from apps.branding.views import brand_dir

    stored = (brand_dir() / BrandAsset.objects.get(pk=resp.data["light"]["id"]).storage_key).read_bytes()
    assert b"<?php" not in stored


def test_brand_settings_reject_markup(world):
    c = client_for(world.admin)
    assert c.patch("/api/v1/branding", {"system_title": "<img src=x onerror=alert(1)>"}, format="json").status_code == 400
    ok = c.patch("/api/v1/branding", {"system_title": "سامانه لیدها", "logo_scale": 1.2}, format="json")
    assert ok.status_code == 200 and ok.data["system_title"] == "سامانه لیدها"
    assert c.patch("/api/v1/branding", {"logo_scale": 5}, format="json").status_code == 400
    assert client_for(world.specialist).patch("/api/v1/branding", {"system_title": "x"}, format="json").status_code == 403
